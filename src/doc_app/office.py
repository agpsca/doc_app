"""Exporta os mesmos blocos para Word (.docx) e PowerPoint (.pptx).

A aparência segue o PDF: mesma paleta, mesmas fontes, mesmas medidas (as do
`Theme`, convertidas de mm para pt), capa escura com faixa de destaque e selo
de classificação, cabeçalho e rodapé com régua, h1 com filete, caixa de
código, callout com barra, tabela de cabeçalho escuro e faixa zebrada,
amostras de cor, e a marca d'água diagonal quando o documento é restrito. No
.pptx, capa e divisórias sangram na cor escura, como no deck em PDF.

O que não dá para igualar (e por quê):

* **Fontes** — o arquivo pede a mesma família do PDF (DejaVu Sans no Linux,
  Arial no Windows). Se a família não estiver instalada na máquina que abre, o
  Office substitui e o desenho muda um pouco.
* **Quebras** — Word e PowerPoint refluem o texto; as quebras de página e de
  linha não caem nos mesmos pontos do PDF, que é medido na geração.
* **Sumário** — no Word é um campo TOC nativo (atualiza sozinho ao abrir), não
  o sumário desenhado do PDF.
* **SVG** — logo e área de proteção em SVG não entram (o Office não lê SVG);
  fica a legenda com o caminho do arquivo.
"""

from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import quoteattr

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn as pqn
from pptx.util import Mm as PptxMm
from pptx.util import Pt as PptxPt

from .dsl import Block
from .renderer import Renderer
from .theme import DEFAULT_THEME, RGB, Theme

# Formatos de imagem que o Office aceita embutido (SVG fica de fora).
RASTER = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}

# Arquivo de fonte do tema -> família que o Office entende.
_FAMILIES = {
    "DejaVuSans.ttf": "DejaVu Sans",
    "DejaVuSans-Bold.ttf": "DejaVu Sans",
    "DejaVuSansMono.ttf": "DejaVu Sans Mono",
    "arial.ttf": "Arial",
    "arialbd.ttf": "Arial",
    "consola.ttf": "Consolas",
}

# Alturas fixas do slide (mm), as mesmas do deck em PDF (slides.py).
LABEL_Y, TITLE_Y, BODY_Y, MAX_LOGO_H = 12.0, 22.0, 46.0, 62.0

PENDING = "A DEFINIR"


def _pt(mm: float) -> float:
    """mm -> pt, para as medidas do tema virarem medidas do Office."""
    return mm * 72 / 25.4


def _line_mm(size_pt: float) -> float:
    """Altura de uma linha de texto nesse corpo, em mm (entrelinha de 1,2)."""
    return size_pt * 1.2 * 25.4 / 72


def _family(filename: str, fallback: str = "Arial") -> str:
    return _FAMILIES.get(filename, fallback)


def _hex(color: RGB) -> str:
    return "{:02X}{:02X}{:02X}".format(*color)


def _hex_rgb(value: str) -> RGB:
    raw = str(value).lstrip("#")
    if len(raw) == 3:
        raw = "".join(c * 2 for c in raw)
    try:
        return int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16)
    except ValueError:
        return (128, 128, 128)


def _asset(spec: dict, theme: Theme, key: str = "path") -> Path | None:
    """Caminho da imagem do bloco, se existir e o Office souber ler."""
    rel = spec.get(key) or ""
    if not rel:
        return None
    path = Path(rel).expanduser()
    if not path.is_absolute():
        if not theme.assets_root:
            return None
        path = Path(theme.assets_root).expanduser() / rel
    return path if path.is_file() and path.suffix.lower() in RASTER else None


def _caption(spec: dict, clearspace: bool = False) -> str:
    if clearspace:
        ratio = float(spec.get("ratio") or 0.5)
        label = "½H" if abs(ratio - 0.5) < 0.01 else f"{ratio:g} H"
        return spec.get("caption") or (
            f"Área de proteção: respiro mínimo de {label} (H é a altura do "
            f"símbolo) nos quatro lados.")
    return spec.get("caption") or spec.get("path", "")


def _no_image(spec: dict) -> str:
    return (f"[imagem não embutida: {spec.get('path', '—')} — o Office não lê "
            f"SVG; veja o PDF]")


def _specimen_text(spec: dict) -> str:
    parts = [spec["font"]]
    if spec.get("role"):
        parts.append(f"papel: {spec['role']}")
    if spec.get("weights"):
        parts.append(f"pesos: {spec['weights']}")
    if spec.get("source"):
        parts.append(f"declarada em {spec['source']}")
    return " · ".join(parts)


def _pending(text: str) -> tuple[bool, str]:
    """Mesma regra do PDF: "(TODO) ..." vira a anotação "A DEFINIR"."""
    return Renderer.pending(str(text))


class _Text:
    """Mede o texto com as mesmas fontes do PDF, para as linhas quebrarem nos
    mesmos pontos. Sem as fontes instaladas, cai numa estimativa pela largura
    média do caractere."""

    def __init__(self, theme: Theme) -> None:
        self.pdf = None
        try:
            from fpdf import FPDF

            pdf = FPDF(format=(theme.page.width, theme.page.height), unit="mm")
            fonts = theme.fonts
            pdf.add_font("sans", "", fonts.path(fonts.sans_regular))
            pdf.add_font("bold", "", fonts.path(fonts.sans_bold))
            pdf.add_font("mono", "", fonts.path(fonts.mono_regular))
            pdf.add_page()
            self.pdf = pdf
        except (FileNotFoundError, RuntimeError):
            self.pdf = None          # sem fontes: segue na estimativa

    def lines(self, text: str, width_mm: float, size: float, *,
              mono: bool = False, bold: bool = False) -> int:
        if self.pdf is None:
            # ponytail: largura média de 0,52 em; só vale como reserva de
            # espaço. Com as fontes do tema instaladas, a medição é exata.
            per_line = max(8, int(width_mm / (size * 0.3528 * 0.52)))
            return max(1, -(-len(text) // per_line))
        self.pdf.set_font("mono" if mono else "bold" if bold else "sans", "", size)
        return max(1, len(self.pdf.multi_cell(width_mm, 5, text, dry_run=True,
                                              output="LINES")))


# --- Word ------------------------------------------------------------------


def build_docx(blocks: list[Block], output_path: str | Path,
               theme: Theme | None = None) -> Path:
    """Gera o .docx em `output_path` a partir dos `blocks`."""
    word = _Word(theme or DEFAULT_THEME)
    for block in blocks:
        word.add(block)
    return word.save(output_path)


def _xml(tag: str, **attrs) -> OxmlElement:
    el = OxmlElement(tag)
    for name, value in attrs.items():
        el.set(qn(f"w:{name}"), str(value))
    return el


def _no_theme(rPr) -> None:
    """Tira os atributos de tema do rPr.

    No OOXML, `w:asciiTheme` e `w:themeColor` vencem `w:ascii` e `w:val`: sem
    removê-los, o Word ignora a fonte e a cor do nosso tema e usa as do
    modelo dele (Calibri Light e azul), que é o que deixava o .docx diferente
    do PDF.
    """
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is not None:
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attr}"), None)
    for tag in ("w:color", "w:shd"):
        node = rPr.find(qn(tag))
        if node is not None:
            for attr in ("themeColor", "themeTint", "themeShade", "themeFill",
                         "themeFillTint", "themeFillShade"):
                node.attrib.pop(qn(f"w:{attr}"), None)


def _shade(element, color: RGB) -> None:
    """Fundo de um parágrafo (w:p) ou de uma célula (w:tc). Não há equivalente
    na API do python-docx; o w:shd entra nas propriedades do nó, que é onde o
    schema do OOXML espera."""
    shd = _xml("w:shd", val="clear", fill=_hex(color))
    props = (element.get_or_add_pPr() if element.tag.endswith("}p")
             else element.get_or_add_tcPr())
    props.append(shd)


def _border(paragraph, side: str, color: RGB, width_mm: float,
            space: int = 0) -> None:
    """Régua/filete num lado do parágrafo. `width_mm` vira oitavos de ponto,
    como o OOXML mede (w:sz)."""
    pPr = paragraph._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    bdr.append(_xml(f"w:{side}", val="single", color=_hex(color),
                    sz=max(2, round(_pt(width_mm) * 8)), space=space))


def _cell_border(cell, side: str, color: RGB, width_mm: float) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    borders.append(_xml(f"w:{side}", val="single", color=_hex(color),
                        sz=max(2, round(_pt(width_mm) * 8))))


class _Word:
    """Monta o .docx com o desenho do PDF: capa escura, corpo com cabeçalho e
    rodapé, e um estilo por tipo de bloco."""

    def __init__(self, theme: Theme) -> None:
        self.t = theme
        self.sans = _family(theme.fonts.sans_regular)
        self.mono = _family(theme.fonts.mono_regular, "Consolas")
        self.doc = Document()
        self._styles()
        self._cover()
        self._body_section()
        self._toc()

    def save(self, output_path: str | Path) -> Path:
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))
        return out

    # -- estrutura ---------------------------------------------------------

    def _styles(self) -> None:
        typo, pal = self.t.typography, self.t.palette
        # o modelo do python-docx traz fonte e cor de tema nos padrões do
        # documento; sem limpar, elas vencem o que definimos aqui
        defaults = self.doc.styles.element.find(qn("w:docDefaults"))
        if defaults is not None:
            for rPr in defaults.iter(qn("w:rPr")):
                _no_theme(rPr)

        normal = self.doc.styles["Normal"]
        normal.font.name = self.sans
        normal.font.size = Pt(typo.body_size)
        normal.font.color.rgb = RGBColor(*pal.ink)
        _no_theme(normal.element.get_or_add_rPr())
        fmt = normal.paragraph_format
        fmt.line_spacing = Pt(_pt(typo.body_leading))
        fmt.space_after = Pt(_pt(2.2))
        fmt.space_before = Pt(0)

        for level, size, color, before, after in (
            (1, typo.h1_size, pal.brand, typo.h1_gap_before, 1.0),
            (2, typo.h2_size, pal.dark, typo.h2_gap_before, 2.0),
            (3, typo.h3_size, pal.ink, 3.0, 1.5),
        ):
            style = self.doc.styles[f"Heading {level}"]
            style.font.name = self.sans
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(*color)
            style.font.italic = False
            _no_theme(style.element.get_or_add_rPr())
            style.paragraph_format.space_before = Pt(_pt(before))
            style.paragraph_format.space_after = Pt(_pt(after))
            style.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number"):
            style = self.doc.styles[name]
            style.font.name = self.sans
            style.font.color.rgb = RGBColor(*pal.ink)
            _no_theme(style.element.get_or_add_rPr())
            # recuo do PDF: marcador na margem, texto 7 mm adiante
            style.paragraph_format.left_indent = Mm(7)
            style.paragraph_format.first_line_indent = Mm(-7)
            style.paragraph_format.space_after = Pt(_pt(0.8))
        # o Word numera e atualiza o sumário sozinho ao abrir o arquivo
        self.doc.settings.element.append(_xml("w:updateFields", val="true"))

    def _run(self, par, text: str = "", *, size: float | None = None,
             color: RGB | None = None, bold: bool = False,
             italic: bool = False, mono: bool = False, leading: float | None = None):
        """Trecho de texto com a fonte e a cor do tema (sem tema do Word)."""
        run = par.add_run(text)
        run.font.name = self.mono if mono else self.sans
        if size is not None:
            run.font.size = Pt(size)
            if leading is None:
                leading = size * 1.2
        run.font.bold, run.font.italic = bold, italic
        if color is not None:
            run.font.color.rgb = RGBColor(*color)
        if leading is not None:
            par.paragraph_format.line_spacing = Pt(leading)
        _no_theme(run._r.get_or_add_rPr())
        return run

    def _cover(self) -> None:
        """Capa: página escura, título grande, faixa de destaque e selo.

        O fundo é uma forma ancorada à página (é como o Word desenha capa) e o
        texto fica numa tabela escura sem borda: se o leitor ignorar a forma,
        o texto continua legível sobre a faixa escura.
        """
        t, pal, cov, typo = self.t, self.t.palette, self.t.cover, self.t.typography
        page = t.page
        section = self.doc.sections[0]
        section.page_width, section.page_height = Mm(page.width), Mm(page.height)
        for side in ("left", "right", "top", "bottom"):
            setattr(section, f"{side}_margin", Mm(0))
        section.header_distance = section.footer_distance = Mm(0)

        anchor = self.doc.add_paragraph()
        anchor.paragraph_format.space_after = Pt(0)
        anchor.paragraph_format.line_spacing = Pt(1)
        anchor.add_run().font.size = Pt(1)
        self._page_rect(anchor, 0, 0, page.width, page.height, pal.dark)

        table = self.doc.add_table(rows=1, cols=1)
        table.autofit = False
        cell = table.rows[0].cells[0]
        cell.width = Mm(page.width)
        self._zero_cell_margins(cell)
        _shade(cell._tc, pal.dark)

        # alturas do PDF (renderer._draw_cover_base), em mm a partir do topo:
        # título em 110, subtítulo em 134, faixa em 150 (96 no development),
        # tagline em 158 e nota em 250. Os espaços abaixo são as diferenças,
        # descontada a altura de cada linha (1,2 × o corpo da fonte).
        accent_y = 150.0 if typo.cover_big else 96.0
        par = cell.paragraphs[0]
        used = 0.0
        if not typo.cover_big:                      # filete do topo (development)
            self._bar(par, pal.accent, 3.0, space_before=0)
            used = 3.0
            par = cell.add_paragraph()
        if t.classification:
            self._badge(par, t.classification)
            used += 18.0 + 9 * 1.2 * 25.4 / 72     # espaço + altura do selo
            par = cell.add_paragraph()

        self._cover_line(par, cov.title, 42, pal.accent, bold=True,
                         space_before=110 - used)
        self._cover_line(cell.add_paragraph(), cov.subtitle, 16, pal.white,
                         space_before=134 - 110 - _line_mm(42))
        self._bar(cell.add_paragraph(), pal.accent, 2.0,
                  space_before=accent_y - 134 - _line_mm(16))
        self._cover_line(cell.add_paragraph(), cov.tagline, 12, pal.cover_text,
                         space_before=158 - accent_y - 2.0)
        self._cover_line(cell.add_paragraph(), cov.footnote, 10,
                         pal.cover_muted, space_before=250 - 158 - _line_mm(12))

    def _zero_cell_margins(self, cell) -> None:
        mar = OxmlElement("w:tcMar")
        for side in ("top", "start", "bottom", "end"):
            mar.append(_xml(f"w:{side}", w=0, type="dxa"))
        cell._tc.get_or_add_tcPr().append(mar)

    def _cover_line(self, par, text: str, size: float, color: RGB,
                    bold: bool = False, space_before: float = 0) -> None:
        par.paragraph_format.space_before = Pt(_pt(max(space_before, 0)))
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.left_indent = Mm(self.t.page.margin)
        par.paragraph_format.right_indent = Mm(self.t.page.margin)
        self._run(par, text or "", size=size, color=color, bold=bold)

    def _bar(self, par, color: RGB, width_mm: float,
             space_before: float = 0) -> None:
        """Faixa horizontal cheia, de ponta a ponta da página."""
        par.paragraph_format.space_before = Pt(_pt(max(space_before, 0)))
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = Pt(1)
        self._run(par, size=1)
        _border(par, "bottom", color, width_mm)

    def _badge(self, par, label: str) -> None:
        """Selo de classificação, com borda, no canto superior direito."""
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        par.paragraph_format.space_before = Pt(_pt(18))
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.right_indent = Mm(self.t.page.margin)
        self._run(par, f"  {label.upper()}  ", size=9, bold=True,
                  color=self.t.palette.accent)
        for side in ("top", "bottom", "left", "right"):
            _border(par, side, self.t.palette.accent, 0.5, space=2)

    def _page_rect(self, paragraph, x: float, y: float, w: float, h: float,
                   color: RGB) -> None:
        """Retângulo ancorado à página, atrás do texto (VML, como o Word usa
        nas capas e marcas d'água dos próprios modelos)."""
        style = (f"position:absolute;margin-left:{_pt(x):.1f}pt;"
                 f"margin-top:{_pt(y):.1f}pt;width:{_pt(w):.1f}pt;"
                 f"height:{_pt(h):.1f}pt;z-index:-251658752;"
                 "mso-position-horizontal-relative:page;"
                 "mso-position-vertical-relative:page")
        paragraph._p.append(parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
            '2006/main" xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:o="urn:schemas-microsoft-com:office:office"><w:pict>'
            f'<v:rect style={quoteattr(style)} fillcolor="#{_hex(color)}" '
            'stroked="f" o:allowincell="f"/></w:pict></w:r>'))

    def _body_section(self) -> None:
        """Seção do corpo: margens do tema, cabeçalho e rodapé com régua."""
        t, page = self.t, self.t.page
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        section.page_width, section.page_height = Mm(page.width), Mm(page.height)
        for side in ("left", "right", "top", "bottom"):
            setattr(section, f"{side}_margin", Mm(page.margin))
        section.header_distance, section.footer_distance = Mm(10), Mm(10)
        self.section = section

        for part, linked in ((section.header, False), (section.footer, False)):
            part.is_linked_to_previous = linked
        self._head_line(section.header.paragraphs[0], t.header_left,
                        t.header_right, rule="bottom")
        self._foot_line(section.footer.paragraphs[0])
        if t.classification:
            self._watermark(section.header.add_paragraph(), t.classification)

    def _tabs(self, par) -> None:
        content = self.t.page.content_width
        tabs = par.paragraph_format.tab_stops
        tabs.add_tab_stop(Mm(content / 2), WD_TAB_ALIGNMENT.CENTER)
        tabs.add_tab_stop(Mm(content), WD_TAB_ALIGNMENT.RIGHT)

    def _small(self, par, text: str, color: RGB, bold: bool = False):
        return self._run(par, text, size=8, color=color, bold=bold, leading=10)

    def _head_line(self, par, left: str, right: str, rule: str) -> None:
        pal = self.t.palette
        self._tabs(par)
        par.paragraph_format.space_after = Pt(_pt(2))
        self._small(par, left, pal.muted)
        self._small(par, f"\t\t{right}", pal.muted)
        _border(par, rule, pal.rule, 0.2, space=4)

    def _foot_line(self, par) -> None:
        """doc_app · CLASSIFICAÇÃO · Página X de Y, com régua em cima."""
        t, pal = self.t, self.t.palette
        self._tabs(par)
        par.paragraph_format.space_before = Pt(_pt(2))
        _border(par, "top", pal.rule, 0.2, space=4)
        self._small(par, t.footer_left, pal.muted)
        if t.classification:
            self._small(par, f"\t{t.classification.upper()} · não distribuir",
                        pal.brand, bold=True)
        else:
            self._small(par, "\t", pal.muted)
        self._small(par, "\tPágina ", pal.muted)
        self._field(par, "PAGE")
        self._small(par, " de ", pal.muted)
        self._field(par, "NUMPAGES")

    def _field(self, par, instr: str, placeholder: str = "1") -> None:
        """Campo do Word (PAGE, NUMPAGES, TOC): o valor é calculado por ele."""
        run = self._small(par, placeholder, self.t.palette.muted)
        field = parse_xml(
            '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/'
            'wordprocessingml/2006/main" w:instr='
            + quoteattr(f" {instr} ") + "/>")
        run._r.addprevious(field)
        field.append(run._r)

    def _watermark(self, par, label: str) -> None:
        """Marca d'água diagonal, 7% de opacidade, como no PDF."""
        pal, page = self.t.palette, self.t.page
        par.paragraph_format.space_after = Pt(0)
        w, h = _pt(page.width * 0.75), _pt(page.height * 0.28)
        style = (f"position:absolute;margin-left:0;margin-top:0;width:{w:.0f}pt;"
                 f"height:{h:.0f}pt;rotation:315;z-index:-251658751;"
                 "mso-position-horizontal:center;"
                 "mso-position-horizontal-relative:margin;"
                 "mso-position-vertical:center;"
                 "mso-position-vertical-relative:margin")
        font = f"font-family:&quot;{self.sans}&quot;;font-size:1pt"
        par._p.append(parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
            '2006/main" xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:o="urn:schemas-microsoft-com:office:office"><w:pict>'
            '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" '
            'adj="10800" path="m@7,0l@8,0m@5,21600l@6,21600e"/>'
            f'<v:shape type="#_x0000_t136" style={quoteattr(style)} '
            f'fillcolor="#{_hex(pal.brand)}" stroked="f" o:allowincell="f">'
            '<v:fill opacity=".07"/>'
            f'<v:textpath style="{font}" string={quoteattr(label.upper())}/>'
            '</v:shape></w:pict></w:r>'))

    def _toc(self) -> None:
        """Sumário: título com filete (como no PDF) e o campo TOC do Word,
        que traz os líderes pontilhados e os números de página."""
        pal, typo = self.t.palette, self.t.typography
        par = self.doc.add_paragraph()
        run = par.add_run("Sumário")
        run.font.name, run.font.size, run.font.bold = self.sans, Pt(typo.h1_size), True
        run.font.color.rgb = RGBColor(*pal.brand)
        par.paragraph_format.space_after = Pt(0)
        rule = self.doc.add_paragraph()
        self._rule(rule, 28.0)
        self._field(self.doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u',
                    "Clique com o botão direito e escolha Atualizar campo "
                    "para montar o sumário.")
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _rule(self, par, width_mm: float) -> None:
        """Filete curto sob um título (h1 e Sumário no PDF)."""
        page = self.t.page
        par.paragraph_format.right_indent = Mm(page.content_width - width_mm)
        par.paragraph_format.space_before = Pt(_pt(1.5))
        par.paragraph_format.space_after = Pt(_pt(5))
        par.paragraph_format.line_spacing = Pt(1)
        par.add_run().font.size = Pt(1)
        _border(par, "bottom", self.t.palette.accent, 0.6)

    # -- blocos ------------------------------------------------------------

    def add(self, block: Block) -> None:
        typ, content = block.type, block.content
        if typ in ("h1", "h2", "h3"):
            self._heading(str(content), int(typ[1]))
        elif typ == "p":
            self._paragraph(str(content))
        elif typ in ("bullet", "num"):
            self._list(content, numbered=(typ == "num"))
        elif typ == "code":
            self._code(str(content))
        elif typ == "callout":
            self._callout(str(content["title"]), str(content["text"]))
        elif typ == "kv":
            self._kv(content)
        elif typ == "table":
            self._table(content)
        elif typ in ("logo", "clearspace"):
            self._image(content, clearspace=(typ == "clearspace"))
        elif typ == "swatches":
            self._swatches(content["items"])
        elif typ == "specimen":
            self._specimen(content)
        elif typ == "spacer":
            par = self.doc.add_paragraph()
            par.paragraph_format.space_after = Pt(_pt(float(content or 0)))
            par.paragraph_format.line_spacing = Pt(1)
            par.add_run().font.size = Pt(1)

    def _heading(self, text: str, level: int) -> None:
        self.doc.add_heading(Renderer._inline_pending(text), level)
        if level == 1:                     # filete de 28 mm, como no PDF
            self._rule(self.doc.add_paragraph(), 28.0)

    def _paragraph(self, text: str):
        open_item, text = _pending(text)
        par = self.doc.add_paragraph()
        if open_item:
            self._pending_mark(par, text)
            return par
        par.add_run(text)
        return par

    def _pending_mark(self, par, text: str) -> None:
        """Anotação de pendência: filete à esquerda, rótulo e o assunto."""
        pal, typo = self.t.palette, self.t.typography
        par.paragraph_format.left_indent = Mm(6)
        label = par.add_run(f"{PENDING} ")
        label.font.name, label.font.bold = self.sans, True
        label.font.size = Pt(max(typo.body_size - 3.0, 7.0))
        label.font.color.rgb = RGBColor(*pal.muted)
        run = par.add_run(text)
        run.font.color.rgb = RGBColor(*pal.muted)
        _border(par, "left", pal.rule, 1.2, space=4)

    def _list(self, items, numbered: bool) -> None:
        pal = self.t.palette
        for item in items:
            open_item, text = _pending(str(item))
            par = self.doc.add_paragraph(
                style="List Number" if numbered else "List Bullet")
            par.paragraph_format.space_after = Pt(_pt(0.8))
            if open_item:
                label = par.add_run(f"{PENDING} ")
                label.font.bold, label.font.size = True, Pt(7.5)
                label.font.color.rgb = RGBColor(*pal.muted)
            run = par.add_run(text)
            if open_item:
                run.font.color.rgb = RGBColor(*pal.muted)

    def _code(self, text: str) -> None:
        """Caixa de código: a largura é a da coluna, como no PDF, e o respiro
        de 3 mm vem do `space` da borda (não de recuo)."""
        pal, typo = self.t.palette, self.t.typography
        par = self.doc.add_paragraph()
        fmt = par.paragraph_format
        fmt.space_before, fmt.space_after = Pt(_pt(1)), Pt(_pt(3))
        fmt.line_spacing = Pt(_pt(5.0))
        lines = text.split("\n")
        for i, line in enumerate(lines):
            run = self._run(par, line, size=typo.mono_size, color=pal.code_fg,
                            mono=True, leading=_pt(5.0))
            if i < len(lines) - 1:
                run.add_break()
        _shade(par._p, pal.code_bg)
        for side in ("top", "bottom", "left", "right"):
            _border(par, side, pal.rule, 0.2, space=round(_pt(3)))

    def _callout(self, title: str, text: str) -> None:
        pal = self.t.palette
        open_item, text = _pending(text)
        if open_item:
            text = f"A definir: {text}"
        par = self.doc.add_paragraph()
        fmt = par.paragraph_format
        fmt.space_before, fmt.space_after = Pt(_pt(1)), Pt(_pt(3))
        self._run(par, f"{title}. ", bold=True, color=pal.dark)
        self._run(par, text, color=pal.ink)
        _shade(par._p, pal.callout_bg)
        # barra de destaque na borda esquerda, com o respiro do PDF
        _border(par, "left", pal.accent, 2.5, space=round(_pt(5.5)))

    def _kv(self, pairs) -> None:
        pal = self.t.palette
        for pair in pairs:
            open_item, value = _pending(str(pair["value"]))
            if open_item:
                value = f"A definir: {value}"
            key = self.doc.add_paragraph()
            key.paragraph_format.space_after = Pt(0)
            self._run(key, str(pair["key"]), size=9.5, color=pal.brand,
                      mono=True, leading=_pt(5.4))
            _border(key, "left", pal.accent, 1.4, space=round(_pt(4)))
            val = self.doc.add_paragraph()
            val.paragraph_format.left_indent = Mm(4)
            val.paragraph_format.space_after = Pt(_pt(2.5))
            self._run(val, value, color=pal.ink, leading=_pt(5.2))

    def _table(self, spec: dict) -> None:
        pal, page = self.t.palette, self.t.page
        header = [Renderer._inline_pending(str(c)) for c in spec.get("header") or []]
        rows = [[Renderer._inline_pending(str(c)) for c in row]
                for row in spec.get("rows") or []]
        cols = len(header) or (len(rows[0]) if rows else 1)
        widths = spec.get("widths") or [page.content_width / cols] * cols
        table = self.doc.add_table(rows=0, cols=cols)
        table.autofit = False
        if header:
            cells = table.add_row().cells
            for cell, text in zip(cells, header):
                _shade(cell._tc, pal.dark)
                self._cell(cell, text, pal.white, bold=True)
        for i, row in enumerate(rows):
            cells = table.add_row().cells
            for cell, text in zip(cells, row):
                if i % 2:
                    _shade(cell._tc, pal.zebra)
                self._cell(cell, text, pal.ink)
                _cell_border(cell, "bottom", pal.rule, 0.2)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Mm(width)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(_pt(4))

    def _cell(self, cell, text: str, color: RGB, bold: bool = False) -> None:
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = Pt(_pt(5.2))
        run = par.add_run(text)
        run.font.name, run.font.size, run.font.bold = self.sans, Pt(9.5), bold
        run.font.color.rgb = RGBColor(*color)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    def _image(self, spec: dict, clearspace: bool) -> None:
        pal = self.t.palette
        path = _asset(spec, self.t)
        if path:
            self.doc.add_paragraph().add_run().add_picture(str(path), height=Mm(48))
        else:
            run = self.doc.add_paragraph().add_run(_no_image(spec))
            run.italic = True
            run.font.color.rgb = RGBColor(*pal.muted)
        par = self.doc.add_paragraph()
        run = par.add_run(_caption(spec, clearspace))
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(*pal.muted)

    def _swatches(self, items: list[dict]) -> None:
        """Amostras em linhas de seis, como no PDF: o quadrado da cor em cima
        e o nome com o código embaixo."""
        pal = self.t.palette
        per_row = 6
        for start in range(0, len(items), per_row):
            row_items = items[start:start + per_row]
            table = self.doc.add_table(rows=2, cols=per_row)
            table.autofit = False
            for i in range(per_row):
                chip, label = table.rows[0].cells[i], table.rows[1].cells[i]
                for cell in (chip, label):
                    cell.width = Mm(self.t.page.content_width / per_row)
                if i >= len(row_items):
                    continue
                item = row_items[i]
                _shade(chip._tc, _hex_rgb(item["hex"]))
                chip.paragraphs[0].paragraph_format.space_after = Pt(0)
                chip.paragraphs[0].add_run().font.size = Pt(_pt(14))
                name = label.paragraphs[0]
                name.paragraph_format.space_after = Pt(0)
                run = name.add_run(str(item.get("name", "")))
                run.font.size, run.font.bold = Pt(8), True
                run.font.color.rgb = RGBColor(*pal.ink)
                code = label.add_paragraph()
                code.paragraph_format.space_after = Pt(0)
                run = code.add_run(str(item.get("hex", "")))
                run.font.name, run.font.size = self.mono, Pt(7.5)
                run.font.color.rgb = RGBColor(*pal.muted)
            self.doc.add_paragraph().paragraph_format.space_after = Pt(_pt(3))

    def _specimen(self, spec: dict) -> None:
        pal = self.t.palette
        par = self.doc.add_paragraph()
        run = par.add_run(_specimen_text(spec))
        run.font.bold, run.font.size = True, Pt(11)
        run.font.color.rgb = RGBColor(*pal.dark)
        note = self.doc.add_paragraph()
        run = note.add_run("O arquivo da fonte não está no repositório; o "
                           "documento usa a fonte do tema.")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(*pal.muted)


# --- PowerPoint ------------------------------------------------------------


def build_pptx(blocks: list[Block], output_path: str | Path,
               theme: Theme | None = None, cover_logo: str | None = None) -> Path:
    """Gera o .pptx em `output_path`, com o desenho do deck em PDF: capa e
    divisórias escuras, um slide por subseção e "(cont.)" no que transborda."""
    theme = theme or DEFAULT_THEME
    deck = _Deck(theme)
    deck.cover(cover_logo)
    agenda = deck.reserve_agenda()
    for block in blocks:
        deck.add(block)
    deck.fill_agenda(agenda)
    return deck.save(output_path)


class _Deck:
    """Monta o deck desenhando por posição (mm), nas mesmas medidas do PDF."""

    def __init__(self, theme: Theme) -> None:
        self.t = theme
        self.sans = _family(theme.fonts.sans_regular)
        self.mono = _family(theme.fonts.mono_regular, "Consolas")
        self.prs = Presentation()
        self.prs.slide_width = PptxMm(theme.page.width)
        self.prs.slide_height = PptxMm(theme.page.height)
        self.layout = self.prs.slide_layouts[6]       # em branco
        self.text = _Text(theme)
        self.section = ""
        self.title = theme.cover.title
        self.y = BODY_Y          # próxima posição livre do corpo, em mm
        self.entries: list[tuple[int, str, int]] = []  # (nível, título, slide)

    def save(self, output_path: str | Path) -> Path:
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out))
        return out

    # -- medidas e formas --------------------------------------------------

    @property
    def page(self):
        return self.t.page

    def _bottom(self) -> float:
        return self.page.height - 20.0

    def _new_slide(self, kind: str):
        slide = self.prs.slides.add_slide(self.layout)
        slide.background.fill.solid()
        dark = kind in ("cover", "divider")
        slide.background.fill.fore_color.rgb = PptxColor(
            *(self.t.palette.dark if dark else self.t.palette.callout_bg))
        self.slide, self.kind = slide, kind
        self.y = BODY_Y
        if kind == "content":
            self._labels()
            self._page_number()
        elif kind == "divider":
            self._page_number()
        return slide

    def _rect(self, x: float, y: float, w: float, h: float, color: RGB,
              radius: bool = False):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            PptxMm(x), PptxMm(y), PptxMm(w), PptxMm(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = PptxColor(*color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        shape.text_frame.text = ""
        return shape

    def _text(self, x: float, y: float, w: float, text: str, *, size: float,
              color: RGB, bold: bool = False, mono: bool = False,
              align=PP_ALIGN.LEFT, h: float = 8.0, leading: float | None = None,
              slide=None):
        box = (slide or self.slide).shapes.add_textbox(
            PptxMm(x), PptxMm(y), PptxMm(w), PptxMm(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        par = frame.paragraphs[0]
        par.alignment = align
        if leading:
            par.line_spacing = PptxPt(_pt(leading))
        run = par.add_run()
        run.text = text
        run.font.size = PptxPt(size)
        run.font.bold = bold
        run.font.name = self.mono if mono else self.sans
        run.font.color.rgb = PptxColor(*color)
        return box

    def _lines(self, text: str, width_mm: float, size: float,
               mono: bool = False, bold: bool = False) -> int:
        return self.text.lines(text, width_mm, size, mono=mono, bold=bold)

    def _advance(self, height_mm: float) -> None:
        self.y += height_mm

    def _room(self, height_mm: float) -> None:
        """Garante espaço no slide atual; senão abre um "(cont.)"."""
        if self.y + height_mm > self._bottom():
            self.content(None)

    # -- slides ------------------------------------------------------------

    def _labels(self) -> None:
        p, pal, t = self.page, self.t.palette, self.t
        self._text(p.margin, LABEL_Y, p.content_width / 2,
                   t.header_left.upper(), size=8, color=pal.muted, h=5)
        self._text(p.margin + p.content_width / 2, LABEL_Y, p.content_width / 2,
                   (self.section or t.header_right).upper(), size=8,
                   color=pal.muted, align=PP_ALIGN.RIGHT, h=5)

    def _page_number(self) -> None:
        p, pal, t = self.page, self.t.palette, self.t
        number = len(self.prs.slides._sldIdLst)
        color = pal.cover_muted if self.kind == "divider" else pal.muted
        if t.classification and self.kind == "content":
            self._text(p.margin, p.height - 12, p.content_width / 2,
                       f"{t.classification.upper()} · não distribuir", size=8,
                       color=pal.brand, bold=True, h=5)
        self._text(p.margin, p.height - 12, p.content_width, str(number),
                   size=8, color=color, align=PP_ALIGN.RIGHT, h=5)

    def _watermark(self) -> None:
        """Marca d'água diagonal translúcida, 7% como no PDF."""
        if not self.t.classification or self.kind != "content":
            return
        p, pal = self.page, self.t.palette
        box = self._text(p.width * 0.1, p.height / 2 - 20, p.width * 0.8,
                         self.t.classification.upper(), size=64, color=pal.brand,
                         bold=True, align=PP_ALIGN.CENTER, h=40)
        box.rotation = 315
        run = box.text_frame.paragraphs[0].runs[0]
        # o alpha vai DENTRO da cor (a:srgbClr), não no a:solidFill: fora
        # dali o PowerPoint ignora e a marca sai opaca
        srgb = run._r.get_or_add_rPr().find(pqn("a:solidFill")).find(
            pqn("a:srgbClr"))
        srgb.append(parse_xml(
            '<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/'
            '2006/main" val="7000"/>'))
        # atrás do conteúdo: move a forma para o começo da árvore
        tree = box._element.getparent()
        tree.remove(box._element)
        tree.insert(2, box._element)

    def cover(self, logo: str | None = None) -> None:
        """Capa escura: logo, título grande, faixa de destaque e selo."""
        p, pal, cov = self.page, self.t.palette, self.t.cover
        self._new_slide("cover")
        y = 52.0
        path = _asset({"path": logo}, self.t) if logo else None
        if path:
            self.slide.shapes.add_picture(str(path), PptxMm(p.margin),
                                          PptxMm(y - 30), height=PptxMm(24))
        self._text(p.margin, y, p.content_width, cov.title, size=40,
                   color=pal.accent, bold=True, h=18)
        self._text(p.margin, y + 24, p.content_width, cov.subtitle, size=16,
                   color=pal.white, h=10)
        self._rect(p.margin, y + 40, 40, 1.6, pal.accent)
        self._text(p.margin, y + 48, p.content_width * 0.7, cov.tagline,
                   size=11, color=pal.cover_text, h=12)
        self._text(p.margin, p.height - 24, p.content_width * 0.7, cov.footnote,
                   size=9, color=pal.cover_muted, h=10)
        if self.t.classification:
            self._badge()

    def _badge(self) -> None:
        p, pal = self.page, self.t.palette
        label = self.t.classification.upper()
        w = max(30.0, len(label) * 2.6 + 8)
        shape = self._rect(p.width - p.margin - w, 18, w, 8, pal.dark)
        shape.line.fill.solid()
        shape.line.fill.fore_color.rgb = PptxColor(*pal.accent)
        shape.line.width = PptxPt(1.5)
        frame = shape.text_frame
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        par = frame.paragraphs[0]
        par.alignment = PP_ALIGN.CENTER
        run = par.add_run()
        run.text = label
        run.font.size, run.font.bold, run.font.name = PptxPt(9), True, self.sans
        run.font.color.rgb = PptxColor(*pal.accent)

    def reserve_agenda(self):
        """Slide de conteúdo ("Conteúdo"), preenchido no fim, quando o número
        de cada seção já é conhecido."""
        slide = self._new_slide("content")
        self._title("Conteúdo")
        return slide

    def divider(self, text: str) -> None:
        p, pal = self.page, self.t.palette
        self.section = text
        self._new_slide("divider")
        self.entries.append((0, text, len(self.prs.slides._sldIdLst)))
        self._text(p.margin, p.height / 2 - 18, p.content_width, text,
                   size=self.t.typography.h1_size, color=pal.accent, bold=True,
                   h=16)
        self._rect(p.margin, p.height / 2 + 14, 40, 1.6, pal.accent)

    def content(self, text: str | None) -> None:
        """Slide de conteúdo. Sem `text`, é a continuação do anterior."""
        if text:
            self.title = text
        elif not self.title.endswith(" (cont.)"):
            self.title += " (cont.)"
        self._new_slide("content")
        if text:
            self.entries.append((1, text, len(self.prs.slides._sldIdLst)))
        self._title(self.title)
        self._watermark()

    def _title(self, text: str) -> None:
        p, pal = self.page, self.t.palette
        self._text(p.margin, TITLE_Y, p.content_width, text,
                   size=self.t.typography.h2_size, color=pal.brand, bold=True,
                   h=10)
        self._rect(p.margin, TITLE_Y + 13, 24, 0.8, pal.accent)

    def fill_agenda(self, slide) -> None:
        """Sumário em duas colunas com o número do slide, como no deck."""
        p, pal = self.page, self.t.palette
        entries = self.entries
        if not entries:
            self._text(p.margin, BODY_Y, p.content_width,
                       "Nenhuma seção ainda.", size=11, color=pal.muted,
                       slide=slide)
            return
        col_w = p.content_width / 2 - 6
        per_col = (len(entries) + 1) // 2
        for i, (level, name, number) in enumerate(entries):
            col, row = divmod(i, per_col) if per_col else (0, 0)
            x = p.margin + col * (col_w + 12)
            y = BODY_Y + row * 10.5
            self._text(x, y, col_w - 10, name,
                       size=13 if level == 0 else 11,
                       color=pal.brand if level == 0 else pal.ink,
                       bold=(level == 0), h=7, slide=slide)
            self._text(x + col_w - 10, y, 10, str(number), size=9,
                       color=pal.muted, align=PP_ALIGN.RIGHT, h=7, slide=slide)

    # -- blocos ------------------------------------------------------------

    def add(self, block: Block) -> None:
        typ, content = block.type, block.content
        pal, typo = self.t.palette, self.t.typography
        if typ == "h1":
            self.divider(str(content))
        elif typ == "h2":
            self.content(str(content))
        elif typ == "h3":
            self._block_text(Renderer._inline_pending(str(content)),
                             size=typo.h3_size, color=pal.dark, bold=True,
                             gap=2.0)
        elif typ == "p":
            open_item, text = _pending(str(content))
            self._block_text((f"{PENDING}  " if open_item else "") + text,
                             size=typo.body_size,
                             color=pal.muted if open_item else pal.ink)
        elif typ in ("bullet", "num"):
            for i, item in enumerate(content, 1):
                open_item, text = _pending(str(item))
                marker = f"{i}." if typ == "num" else "•"
                self._bullet(marker, text, open_item)
        elif typ == "code":
            self._code(str(content))
        elif typ == "callout":
            self._callout(str(content["title"]), str(content["text"]))
        elif typ == "kv":
            self._kv(content)
        elif typ == "table":
            self._table(content)
        elif typ == "logo":
            self._logo(content)
        elif typ == "clearspace":
            self._clearspace(content)
        elif typ == "swatches":
            self._swatches(content["items"])
        elif typ == "specimen":
            self._specimen(content)
        elif typ == "spacer":
            self._advance(float(content or 0))

    def _block_text(self, text: str, *, size: float, color: RGB,
                    bold: bool = False, mono: bool = False, indent: float = 0.0,
                    gap: float = 2.2) -> None:
        width = self.page.content_width - indent
        leading = self.t.typography.body_leading
        height = self._lines(text, width, size) * leading + 1
        self._room(height)
        self._text(self.page.margin + indent, self.y, width, text, size=size,
                   color=color, bold=bold, mono=mono, h=height, leading=leading)
        self._advance(height + gap)

    def _bullet(self, marker: str, text: str, open_item: bool) -> None:
        pal = self.t.palette
        size = self.t.typography.body_size
        width = self.page.content_width - 7
        leading = self.t.typography.body_leading
        height = self._lines(text, width, size) * leading + 1
        self._room(height)
        self._text(self.page.margin, self.y, 7, "·" if open_item else marker,
                   size=size, color=pal.muted if open_item else pal.ink, h=height)
        self._text(self.page.margin + 7, self.y, width,
                   (f"{PENDING}  " if open_item else "") + text, size=size,
                   color=pal.muted if open_item else pal.ink, h=height,
                   leading=leading)
        self._advance(height + 0.8)

    def _code(self, text: str) -> None:
        pal, typo = self.t.palette, self.t.typography
        lines = text.split("\n")
        line_h, pad = 5.0, 3.0
        height = len(lines) * line_h + 2 * pad
        self._room(min(height, 40))
        height = min(height, self._bottom() - self.y)
        box = self._rect(self.page.margin, self.y, self.page.content_width,
                         height, pal.code_bg)
        box.line.fill.solid()
        box.line.fill.fore_color.rgb = PptxColor(*pal.rule)
        box.line.width = PptxPt(0.5)
        frame = box.text_frame
        frame.margin_left = frame.margin_top = PptxMm(pad)
        frame.word_wrap = False
        for i, line in enumerate(lines):
            par = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            par.line_spacing = PptxPt(_pt(line_h))
            run = par.add_run()
            run.text = line
            run.font.size, run.font.name = PptxPt(typo.mono_size), self.mono
            run.font.color.rgb = PptxColor(*pal.code_fg)
        self._advance(height + 3)

    def _callout(self, title: str, text: str) -> None:
        pal = self.t.palette
        open_item, text = _pending(text)
        if open_item:
            text = f"A definir: {text}"
        full = f"{title}. {text}"
        leading = self.t.typography.body_leading
        inner = self.page.content_width - 10
        height = self._lines(full, inner, self.t.typography.body_size) * leading + 6
        self._room(height)
        self._rect(self.page.margin, self.y, self.page.content_width, height,
                   pal.callout_bg)
        self._rect(self.page.margin, self.y, 2.5, height, pal.accent)
        box = self._text(self.page.margin + 7, self.y + 3, inner, f"{title}. ",
                         size=self.t.typography.body_size, color=pal.dark,
                         bold=True, h=height - 6, leading=leading)
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = PptxPt(self.t.typography.body_size)
        run.font.name = self.sans
        run.font.color.rgb = PptxColor(*pal.ink)
        self._advance(height + 3)

    def _kv(self, pairs) -> None:
        pal = self.t.palette
        for pair in pairs:
            open_item, value = _pending(str(pair["value"]))
            if open_item:
                value = f"A definir: {value}"
            self._room(12)
            self._rect(self.page.margin, self.y + 0.6, 1.4, 4.2, pal.accent)
            self._block_text(str(pair["key"]), size=9.5, color=pal.brand,
                             mono=True, indent=4, gap=0.4)
            self._block_text(value, size=10, color=pal.ink, indent=4, gap=2.5)

    def _table(self, spec: dict) -> None:
        pal, p = self.t.palette, self.page
        header = [Renderer._inline_pending(str(c)) for c in spec.get("header") or []]
        rows = [[Renderer._inline_pending(str(c)) for c in row]
                for row in spec.get("rows") or []]
        data = ([header] if header else []) + rows
        if not data:
            return
        cols = max(len(r) for r in data)
        row_h = 9.0
        height = row_h * len(data)
        self._room(min(height, 60))
        height = min(height, self._bottom() - self.y)
        shape = self.slide.shapes.add_table(
            len(data), cols, PptxMm(p.margin), PptxMm(self.y),
            PptxMm(p.content_width), PptxMm(height))
        table = shape.table
        table.first_row = bool(header)
        table.horz_banding = False
        widths = spec.get("widths") or [p.content_width / cols] * cols
        for c in range(cols):
            table.columns[c].width = PptxMm(widths[c] if c < len(widths)
                                            else p.content_width / cols)
        for r, row in enumerate(data):
            head = bool(header) and r == 0
            zebra = not head and (r - (1 if header else 0)) % 2
            for c in range(cols):
                cell = table.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = PptxColor(
                    *(pal.dark if head else pal.zebra if zebra else pal.white))
                cell.margin_left = cell.margin_right = PptxMm(2)
                cell.margin_top = cell.margin_bottom = PptxMm(1)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                par = cell.text_frame.paragraphs[0]
                run = par.add_run()
                run.text = row[c] if c < len(row) else ""
                run.font.size, run.font.name = PptxPt(9.5), self.sans
                run.font.bold = head
                run.font.color.rgb = PptxColor(*(pal.white if head else pal.ink))
        self._advance(height + 4)

    def _logo(self, spec: dict) -> None:
        """Asset da marca. Com `on: both`, a mesma marca sobre fundo claro e
        sobre fundo escuro, como no deck em PDF."""
        pal, p = self.t.palette, self.page
        top = self.y
        height = min(MAX_LOGO_H, self._bottom() - 10 - top)
        if height < 24:
            self.content(None)
            top, height = self.y, MAX_LOGO_H
        if spec.get("on") == "both":
            half = (p.content_width - 8) / 2
            light = self._rect(p.margin, top, half, height, pal.white)
            light.line.fill.solid()
            light.line.fill.fore_color.rgb = PptxColor(*pal.rule)
            self._rect(p.margin + half + 8, top, half, height, pal.dark)
            self._place(spec, p.margin + 6, top + 8, half - 12, height - 22)
            self._place(spec, p.margin + half + 14, top + 8, half - 12,
                        height - 22, key="path_dark")
            for x, label, color in ((p.margin, "SOBRE FUNDO CLARO", pal.muted),
                                    (p.margin + half + 8, "SOBRE FUNDO ESCURO",
                                     pal.cover_muted)):
                self._text(x + 2, top + height - 9, half, label, size=8,
                           color=color, h=5)
        else:
            if spec.get("on") == "dark":
                self._rect(p.margin, top, p.content_width, height, pal.dark)
            self._place(spec, p.margin + 8, top + 6, p.content_width - 16,
                        height - 12)
        self.y = top + height + 1
        self._block_text(_caption(spec), size=9, color=pal.muted, gap=2)

    def _place(self, spec: dict, x: float, y: float, w: float, h: float,
               key: str = "path") -> None:
        path = _asset(spec, self.t, key) or _asset(spec, self.t)
        if path:
            self.slide.shapes.add_picture(str(path), PptxMm(x), PptxMm(y),
                                          height=PptxMm(h))
            return
        self._text(x, y + h / 2 - 3, w, _no_image(spec), size=8,
                   color=self.t.palette.muted, align=PP_ALIGN.CENTER, h=10)

    def _clearspace(self, spec: dict) -> None:
        """Área de proteção: grade de ¼H com o símbolo no centro."""
        pal, p = self.t.palette, self.page
        ratio = float(spec.get("ratio") or 0.5)
        H = 36.0
        pad = H * ratio
        box = H + 2 * pad
        module = H / 4
        self._room(box + 2 * module + 12)
        top = self.y + 4
        x0 = p.margin + (p.content_width - box) / 2
        gx, gy = x0 - module, top - module
        steps = int((box + 2 * module) / module)
        for i in range(steps + 1):          # grade discreta atrás do símbolo
            self._rect(gx + i * module, gy, 0.1, steps * module, pal.rule)
            self._rect(gx, gy + i * module, steps * module, 0.1, pal.rule)
        self._place(spec, x0 + pad, top + pad, H, H)
        self.y = top + box + module + 2
        self._block_text(_caption(spec, clearspace=True), size=9,
                         color=pal.muted)

    def _swatches(self, items: list[dict]) -> None:
        pal, p = self.t.palette, self.page
        per_row = 6
        chip = (p.content_width - (per_row - 1) * 3) / per_row
        for start in range(0, len(items), per_row):
            row = items[start:start + per_row]
            self._room(chip + 14)
            top = self.y
            for i, item in enumerate(row):
                x = p.margin + i * (chip + 3)
                self._rect(x, top, chip, chip, _hex_rgb(item["hex"]), radius=True)
                self._text(x, top + chip + 1, chip, str(item.get("name", "")),
                           size=8, color=pal.ink, bold=True, h=4)
                self._text(x, top + chip + 5, chip, str(item.get("hex", "")),
                           size=7.5, color=pal.muted, mono=True, h=4)
            self._advance(chip + 14)

    def _specimen(self, spec: dict) -> None:
        pal = self.t.palette
        self._block_text(_specimen_text(spec), size=11, color=pal.dark,
                         bold=True, gap=0.5)
        self._block_text("O arquivo da fonte não está no repositório; o deck "
                         "usa a fonte do tema.", size=9, color=pal.muted)
