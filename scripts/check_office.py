#!/usr/bin/env python
"""Checagem da exportação para Word (.docx) e PowerPoint (.pptx).

Monta um documento com todos os tipos de bloco da DSL, exporta nos dois
formatos e confere duas coisas:

1. **Abre**: pacote OOXML válido (zip íntegro, partes obrigatórias, todo XML
   bem formado) e relido pelas bibliotecas. É o que o Office recusa como
   "conteúdo ilegível".
2. **Parece com o PDF**: capa escura com faixa e selo, cabeçalho e rodapé com
   régua e número de página, sumário, filete do h1, caixa de código, callout
   com barra, tabela de cabeçalho escuro com zebra, amostras na cor certa,
   marca d'água diagonal, e as fontes e medidas do tema.

Uso:
    uv run python scripts/check_office.py
"""

from __future__ import annotations

import base64
import dataclasses
import re
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from docx import Document
from pptx import Presentation

from doc_app import office
from doc_app.dsl import BLOCK_TYPES, Block
from doc_app.office import build_docx, build_pptx
from doc_app.theme import BRAND_THEME, SLIDES_THEME

# PNG 1x1 válido: só para provar que imagem em formato raster entra no arquivo.
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFAAH/"
    "q842iQAAAABJRU5ErkJggg==")

MARK = "Uso interno"


def blocks() -> list[Block]:
    """Um bloco de cada tipo da DSL, para nenhum ficar sem tradução."""
    doc = [
        Block("h1", "1. A marca"),
        Block("p", "Parágrafo de abertura."),
        Block("h2", "1.1. Regras"),
        Block("p", "(TODO) O que ainda falta decidir."),
        Block("bullet", ["primeiro item", "segundo item"]),
        Block("num", ["passo um", "passo dois"]),
        Block("code", "def soma(a, b):\n    return a + b"),
        Block("callout", {"title": "Atenção", "text": "Não mexa no YAML gerado."}),
        Block("kv", [{"key": "Stack", "value": "Python"}]),
        Block("table", {"header": ["Arquivo", "Papel"],
                        "rows": [["cli.py", "painel"], ["core.py", "operações"],
                                 ["dsl.py", "blocos"]]}),
        Block("spacer", 4),
        Block("h3", "Subtítulo"),
        Block("logo", {"path": "logo.png", "caption": "Assinatura principal"}),
        Block("clearspace", {"path": "logo.svg", "ratio": 0.5}),
        Block("swatches", {"items": [{"name": "Principal", "hex": "#5168DA"},
                                     {"name": "Escura", "hex": "#242F62"}]}),
        Block("specimen", {"font": "Poppins", "role": "títulos",
                           "weights": "400, 600"}),
    ]
    faltando = BLOCK_TYPES - {b.type for b in doc}
    assert not faltando, f"tipos de bloco sem checagem: {sorted(faltando)}"
    return doc


def parts(path: Path) -> dict[str, bytes]:
    """Confere o pacote OOXML e devolve as partes, para inspeção."""
    assert path.is_file() and path.stat().st_size > 0, path
    with zipfile.ZipFile(path) as zf:
        assert zf.testzip() is None, f"zip corrompido: {path}"
        content = {name: zf.read(name) for name in zf.namelist()}
    for name in ("[Content_Types].xml", "_rels/.rels"):
        assert name in content, f"falta {name} em {path.name}"
    for name, data in content.items():
        if name.endswith((".xml", ".rels")):
            ElementTree.fromstring(data)      # levanta se estiver mal formado
    return content


def check_docx(assets: Path, out: Path) -> None:
    theme = dataclasses.replace(BRAND_THEME, assets_root=str(assets),
                                classification=MARK)
    pal, sans = theme.palette, office._family(theme.fonts.sans_regular)
    mono = office._family(theme.fonts.mono_regular, "Consolas")
    path = build_docx(blocks(), out / "doc.docx", theme)
    content = parts(path)
    assert "word/document.xml" in content
    xml = content["word/document.xml"].decode()
    head = "".join(v.decode() for k, v in content.items()
                   if k.startswith(("word/header", "word/footer")))

    doc = Document(str(path))
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    # todo o texto, inclusive o de dentro das tabelas (a capa é uma tabela)
    texto = "\n".join(
        node.text or "" for node in
        ElementTree.fromstring(content["word/document.xml"]).iter(W + "t"))

    # --- página e fontes do tema (o Word grava em twips: tolerância de 0,1 mm)
    capa, corpo = doc.sections[0], doc.sections[1]
    assert abs(capa.page_width.mm - theme.page.width) < 0.1, capa.page_width.mm
    assert capa.left_margin.mm == 0, "a capa sangra até a borda"
    assert abs(corpo.left_margin.mm - theme.page.margin) < 0.1, corpo.left_margin.mm
    assert doc.styles["Normal"].font.name == sans

    # --- capa: fundo escuro, faixa de destaque, selo e textos
    assert f'fillcolor="#{office._hex(pal.dark)}"' in xml, "capa sem fundo escuro"
    assert office._hex(pal.dark) in xml and office._hex(pal.accent) in xml
    for linha in (theme.cover.title, theme.cover.subtitle, theme.cover.tagline,
                  theme.cover.footnote):
        assert linha in texto, f"capa sem {linha!r}"
    assert MARK.upper() in texto, "capa sem o selo de classificação"

    # --- cabeçalho, rodapé, sumário e marca d'água
    assert theme.header_left in doc.sections[1].header.paragraphs[0].text
    assert theme.footer_left in doc.sections[1].footer.paragraphs[0].text
    assert "w:instr" in xml and "TOC" in xml, "falta o campo de sumário"
    assert "PAGE" in head and "NUMPAGES" in head, "rodapé sem número de página"
    assert "v:textpath" in head and MARK.upper() in head, "falta marca d'água"
    assert 'opacity=".07"' in head, "marca d'água sem transparência"

    # --- blocos
    estilos = [p.style.name for p in doc.paragraphs]
    for level in (1, 2, 3):
        assert f"Heading {level}" in estilos, estilos      # sumário do Word
    assert "List Bullet" in estilos and "List Number" in estilos, estilos
    assert doc.styles["Heading 1"].font.color.rgb == office.RGBColor(*pal.brand)
    for nome in ("Normal", "Heading 1", "Heading 2", "Heading 3",
                 "List Bullet", "List Number"):
        estilo = doc.styles[nome].element.xml
        for atributo in ("themeColor", "asciiTheme", "hAnsiTheme", "themeShade"):
            assert atributo not in estilo, (
                f"{nome} com {atributo}: no OOXML o tema vence a cor e a fonte "
                f"que definimos, e o Word usaria as do modelo dele")
    assert any(r.font.name == mono for p in doc.paragraphs for r in p.runs), \
        "bloco de código sem a fonte monoespaçada do tema"
    assert office.PENDING in texto, "parágrafo (TODO) sem a anotação A DEFINIR"
    assert "Poppins · papel: títulos · pesos: 400, 600" in texto, texto
    assert "logo.svg" in texto and "não lê SVG" in texto, "falta o aviso do SVG"
    assert len(doc.inline_shapes) == 1, "PNG deveria entrar como imagem"

    for nome, cor in (("código", pal.code_bg), ("callout", pal.callout_bg),
                      ("cabeçalho da tabela", pal.dark),
                      ("zebra da tabela", pal.zebra)):
        assert f'w:fill="{office._hex(cor)}"' in xml, f"fundo do {nome} faltando"
    assert 'w:fill="5168DA"' in xml, "amostra de cor não foi pintada"

    # tabela com as larguras pedidas e o cabeçalho em negrito branco
    tabela = next(t for t in doc.tables if len(t.columns) == 2 and len(t.rows) == 4)
    cab = tabela.rows[0].cells[0].paragraphs[0].runs[0]
    assert cab.font.bold and cab.font.color.rgb == office.RGBColor(*pal.white)


def check_pptx(assets: Path, out: Path) -> None:
    theme = dataclasses.replace(SLIDES_THEME, assets_root=str(assets),
                                classification=MARK)
    pal = theme.palette
    path = build_pptx(blocks(), out / "deck.pptx", theme, cover_logo="logo.png")
    content = parts(path)
    assert "ppt/presentation.xml" in content and "ppt/slides/slide1.xml" in content
    xml = "".join(v.decode() for k, v in content.items()
                  if k.startswith("ppt/slides/slide"))

    prs = Presentation(str(path))
    assert prs.slide_width == 12192120 and prs.slide_height == 6858000, \
        (prs.slide_width, prs.slide_height)                       # 16:9
    slides = list(prs.slides)
    assert len(slides) >= 5, len(slides)

    def fundo(slide):
        return slide.background.fill.fore_color.rgb

    assert fundo(slides[0]) == office.PptxColor(*pal.dark), "capa não é escura"
    assert fundo(slides[1]) == office.PptxColor(*pal.callout_bg), \
        "slide de conteúdo não tem o fundo claro do deck"

    def textos(slide):
        return "\n".join(sh.text_frame.text for sh in slide.shapes
                         if sh.has_text_frame)

    capa = textos(slides[0])
    for linha in (theme.cover.title, theme.cover.subtitle, theme.cover.tagline):
        assert linha in capa, f"capa sem {linha!r}"
    assert MARK.upper() in capa, "capa sem o selo"
    assert any(sh.shape_type is not None and "PICTURE" in str(sh.shape_type)
               for sh in slides[0].shapes), "logo da capa não entrou"

    assert "Conteúdo" in textos(slides[1]), "falta o slide de sumário"
    assert "1. A marca" in textos(slides[1]), "sumário sem as seções"

    todos = "\n".join(textos(s) for s in slides)
    assert theme.header_left.upper() in todos, "slides sem a etiqueta do topo"
    assert "1. A marca" in todos and "1.1. Regras" in todos
    assert "(cont.)" in todos, "nada transbordou para um slide de continuação"
    assert office.PENDING in todos, "bloco (TODO) sem a anotação A DEFINIR"
    assert "não lê SVG" in todos, "falta o aviso do SVG"
    assert f"{MARK.upper()} · não distribuir" in todos, "slides sem a restrição"

    tipos = [str(sh.shape_type) for s in slides for sh in s.shapes]
    assert any("TABLE" in t for t in tipos), tipos
    assert any("PICTURE" in t for t in tipos), "PNG deveria entrar como imagem"
    assert 'prst="roundRect"' in xml, "amostras de cor não entraram"
    # o alpha precisa estar DENTRO da cor, senão o PowerPoint ignora
    assert re.search(r'<a:srgbClr val="[0-9A-F]{6}"><a:alpha val="7000"/>', xml), \
        "marca d'água opaca (alpha fora do a:srgbClr)"
    assert office._hex(pal.accent) in xml, "filetes de destaque faltando"
    assert "5168DA" in xml, "amostra de cor não foi pintada"

    # nada pode passar da borda do slide (a marca d'água é girada de propósito)
    for i, slide in enumerate(slides):
        for sh in slide.shapes:
            if sh.rotation:
                continue
            assert (sh.top or 0) >= 0 and (sh.left or 0) >= 0, f"slide {i}"
            assert (sh.top or 0) + (sh.height or 0) <= prs.slide_height + 1, \
                f"forma passa da borda de baixo no slide {i}"
            assert (sh.left or 0) + (sh.width or 0) <= prs.slide_width + 1, \
                f"forma passa da borda direita no slide {i}"


def main() -> int:
    with tempfile.TemporaryDirectory() as raw:
        tmp = Path(raw)
        assets = tmp / "app"
        assets.mkdir()
        (assets / "logo.png").write_bytes(PNG)
        (assets / "logo.svg").write_text("<svg/>", encoding="utf-8")
        check_docx(assets, tmp)
        check_pptx(assets, tmp)
    print("ok: docx e pptx válidos, com a capa, as cores, as fontes e a marca "
          "do PDF")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
